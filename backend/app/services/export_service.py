import io
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List
from backend.app.models.property_model import Property

def export_properties_to_excel(properties: List[Property]) -> io.BytesIO:
    data = []
    now = datetime.utcnow()
    for p in properties:
        days_pub = (now - p.publication_date).days if p.publication_date else None
        data.append({
            'ID': p.id,
            'Portal': p.portal.capitalize(),
            'Tipo': p.property_type.capitalize(),
            'Título': p.title,
            'Precio USD': p.price_usd,
            'Superficie Total (m²)': p.total_area_m2,
            'Superficie Cubierta (m²)': p.covered_area_m2,
            'Precio USD/m²': round(p.price_per_m2, 1) if p.price_per_m2 else None,
            'Dormitorios': p.bedrooms,
            'Baños': p.bathrooms,
            'Ambientes': p.rooms,
            'Barrio': p.neighborhood,
            'Zona': p.zone,
            'Dirección': p.address,
            'Expensas': p.expenses,
            'Inmobiliaria': p.seller_name,
            'Antigüedad': p.antiquity or 'N/D',
            'Disposición': p.disposition or 'N/D',
            'Orientación': p.orientation or 'N/D',
            'Días Publicado': f"{days_pub} días" if days_pub is not None else 'N/D',
            'Fecha Publicación': p.publication_date.strftime('%d/%m/%Y') if p.publication_date else 'N/D',
            'Enlace Original': p.url,
            'Es Favorito': '⭐ Sí' if p.is_favorite else 'No',
            'Grupo Duplicado': p.cluster_id or 'Único',
            'Notas': p.user_notes or ''
        })
        
    df = pd.DataFrame(data)
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Inmuebles')
        worksheet = writer.sheets['Inmuebles']
        
        # Styling
        header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell_font = Font(name='Arial', size=10)
        link_font = Font(name='Arial', size=10, color='2563EB', underline='single')
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        
        # Style Header
        for col_num, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
        # Style Data Rows
        url_col_idx = list(df.columns).index('Enlace Original') + 1
        price_col_idx = list(df.columns).index('Precio USD') + 1
        m2_price_col_idx = list(df.columns).index('Precio USD/m²') + 1
        
        for row_num in range(2, len(df) + 2):
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = cell_font
                cell.border = thin_border
                
                # Format URL as clickable hyperlink
                if col_num == url_col_idx:
                    url_val = cell.value
                    if url_val and str(url_val).startswith('http'):
                        cell.hyperlink = url_val
                        cell.font = link_font
                        
                # Format Currency
                if col_num in (price_col_idx, m2_price_col_idx) and isinstance(cell.value, (int, float)):
                    cell.number_format = '"$"#,##0'
                    
        # Adjust Column Widths
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
            
    output.seek(0)
    return output
